# common/views/import_excel.py
"""
Generic Excel Import engine.

Usage — register a type from any module:
    from common.views.import_excel import register_import_type
    register_import_type('employee', { label, table, pk_col,
                                       db_alias_fn, columns })

columns format:
    [{'excel': 'Reg No', 'db_col': 'RegNo', 'type': 'str', 'required': True}, ...]
    type: 'str' | 'int' | 'float' | 'date'  (date → stored as YYYY-MM-DD)

    ItemGroup lookup columns additionally carry:
        'itemgroup_col': True
        'category_id':  <int>   — ItemGroups.Category value for that field
    The engine resolves the text → GroupID by case-insensitive match inside
    ItemGroups for the given category; if no match exists a new row is
    inserted automatically, then the integer GroupID is stored in the target
    column.  This lookup is done once per unique (category_id, description)
    pair per request and cached for the lifetime of that request.

Template download
    GET /import/template/?type=<key>
    Returns a formatted .xlsx file whose first row contains the column headers
    for the requested import type.  Works for every registered type.
"""
import io
import json
import logging
from datetime import datetime, date, timedelta

from django.db import connections
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.views.decorators.http import require_http_methods

from common.theme_constants import tb, TOOLBAR_ICONS

logger = logging.getLogger(__name__)

# ── Registry ──────────────────────────────────────────────────────────────────
_IMPORT_REGISTRY: dict = {}


def register_import_type(key: str, config: dict):
    """Idempotent — safe to call multiple times."""
    _IMPORT_REGISTRY[key] = config


def get_import_types() -> dict:
    return _IMPORT_REGISTRY


# ── form_config helper ────────────────────────────────────────────────────────

def build_import_form_config(import_types: list) -> dict:
    """
    Returns a profile_form-compatible form_config dict for the import page.

    Parameters
    ----------
    import_types : list
        [{'key': 'employee', 'label': 'Employee'}, ...]
        Typically produced by:
            [{'key': k, 'label': v['label']} for k, v in get_import_types().items()]
    """
    return {
        'form_id': 'ix-form',

        # ── Hero card ──────────────────────────────────────────────────────
        'hero': {
            'show_avatar':    False,
            'hero_icon':      TOOLBAR_ICONS.get('Import', '📥'),
            'default_name':   'Import from Excel',
            'show_meta_rows': False,
            'badge1_label':   '',
            'badge2_field':   '',
        },

        # ── Zone 3 — header fields ─────────────────────────────────────────
        'header_fields': [
            {
                'name':     'ix_import_type',
                'label':    'Import Type',
                'type':     'select',
                'required': True,
                'options':  [{'value': t['key'], 'label': t['label']} for t in import_types],
                'onchange': 'ixTypeChanged()',
                'width':    '200px',
            },
            {
                'name':          'ix_file_display',
                'label':         'Excel / CSV File',
                'type':          'text',
                'readonly':      True,
                'placeholder':   'No file chosen',
                'browse_btn':    True,
                'browse_label':  'Browse…',
                'browse_accept': '.xlsx,.xls,.csv',
                'width':         '280px',
            },
        ],

        # ── Toolbar ────────────────────────────────────────────────────────
        # Template button sits between Edit and Import so the flow reads:
        # "select all → edit → download blank template → import"
        'toolbar': [
            tb('Select All',   'ixSelectAll(true)',    icon=TOOLBAR_ICONS.get('Approve')),
            tb('Deselect All', 'ixSelectAll(false)',   icon=TOOLBAR_ICONS.get('Remove')),
            tb('Edit',         'ixToggleEdit()',       icon=TOOLBAR_ICONS.get('Edit')),
            tb('Template',     'ixDownloadTemplate()', icon=TOOLBAR_ICONS.get('Export', '📄')),
            tb('Import',       'ixImport()',           icon=TOOLBAR_ICONS.get('Import')),
        ],

        'show_sidenav': False,
        'tabs': [],
    }


# ── Import page view ──────────────────────────────────────────────────────────

def import_excel_view(request):
    """
    Renders the generic import-from-Excel page.
    App wrappers (hrms/views/import_view.py, etc.) call this logic themselves
    via build_import_form_config(); this view is used when the common URL
    /import/ is hit directly.
    """
    import_types = [
        {'key': k, 'label': v.get('label', k)}
        for k, v in _IMPORT_REGISTRY.items()
    ]

    columns_config = {
        k: v.get('columns', [])
        for k, v in _IMPORT_REGISTRY.items()
    }

    return render(request, 'common/masters/import_excel.html', {
        'form_config':    build_import_form_config(import_types),
        'import_types':   import_types,
        'columns_config': json.dumps(columns_config),
        'process_url':    'process/',
        'template_url':   'template/',
    })


# ── Template download ─────────────────────────────────────────────────────────

def import_template_download(request):
    """
    GET  /import/template/?type=<key>        (common URL)
    GET  /hrms/import/template/?type=<key>   (app-level URL, same view)

    Generates and streams a .xlsx file whose first row contains the column
    headers for the requested import type.  Works for every registered type.
    The second row contains a short hint in each cell so the user knows what
    format is expected (date columns show DD/MM/YYYY, required columns are
    marked with *).
    """
    # Lazy import — openpyxl is only needed here, not on every request
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return HttpResponse(
            'openpyxl is required for template download. '
            'Run: pip install openpyxl',
            status=500,
            content_type='text/plain',
        )

    # Allow any registered module to call register first (idempotent)
    type_key = request.GET.get('type', '').strip()
    if not type_key or type_key not in _IMPORT_REGISTRY:
        return HttpResponse(
            f'Unknown import type: "{type_key}". '
            f'Available: {", ".join(_IMPORT_REGISTRY.keys())}',
            status=400,
            content_type='text/plain',
        )

    config  = _IMPORT_REGISTRY[type_key]
    columns = config.get('columns', [])
    label   = config.get('label', type_key)

    wb = Workbook()
    ws = wb.active
    ws.title = label

    # ── Styles ────────────────────────────────────────────────────────────────
    header_font    = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    header_fill    = PatternFill('solid', start_color='2563EB', end_color='2563EB')  # blue
    req_fill       = PatternFill('solid', start_color='1D4ED8', end_color='1D4ED8')  # darker blue for required
    hint_font      = Font(name='Arial', italic=True, color='6B7280', size=9)
    hint_fill      = PatternFill('solid', start_color='F3F4F6', end_color='F3F4F6')
    center_align   = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align     = Alignment(horizontal='left',   vertical='center', wrap_text=True)
    thin_side      = Side(style='thin', color='CBD5E1')
    thin_border    = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # ── Row 1: Headers ────────────────────────────────────────────────────────
    for col_idx, col in enumerate(columns, start=1):
        header_text = col['excel']
        if col.get('required'):
            header_text += ' *'

        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.font      = header_font
        cell.fill      = req_fill if col.get('required') else header_fill
        cell.alignment = center_align
        cell.border    = thin_border

    # ── Row 2: Hints (format / example) ──────────────────────────────────────
    _type_hints = {
        'str':   'Text',
        'int':   'Number',
        'float': '0.00',
        'date':  'DD/MM/YYYY',
    }
    for col_idx, col in enumerate(columns, start=1):
        hint = _type_hints.get(col.get('type', 'str'), 'Text')
        # ItemGroup fields — user enters the name, not a number
        if col.get('itemgroup_col'):
            hint = 'Name (text)'
        cell = ws.cell(row=2, column=col_idx, value=hint)
        cell.font      = hint_font
        cell.fill      = hint_fill
        cell.alignment = center_align
        cell.border    = thin_border

    # ── Column widths ─────────────────────────────────────────────────────────
    for col_idx, col in enumerate(columns, start=1):
        # Width = longest of header text, hint text, plus a little padding
        header_len = len(col['excel']) + (2 if col.get('required') else 0)
        hint_len   = len(_type_hints.get(col.get('type', 'str'), 'Text'))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(header_len, hint_len, 12) + 2

    # ── Row heights ───────────────────────────────────────────────────────────
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 16

    # ── Freeze pane below header ──────────────────────────────────────────────
    ws.freeze_panes = 'A3'

    # ── Stream response ───────────────────────────────────────────────────────
    buf      = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f'{type_key}_import_template.xlsx'
    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# ── ItemGroup resolver ────────────────────────────────────────────────────────

def _resolve_itemgroup_id(db_alias: str, category_id: int, description: str,
                          cache: dict) -> int:
    """
    Return the GroupID for *description* inside *category_id* in ItemGroups.

    Lookup is case-insensitive.  If no match is found a new row is inserted
    (GroupID = MAX(GroupID) + 1 across the whole table) and that new ID is
    returned.

    *cache* is a plain dict scoped to the current import request — keyed by
    (category_id, description.lower()) — so repeated occurrences of the same
    value don't hit the DB more than once.

    Parameters
    ----------
    db_alias    : Django database alias ('customer_db', etc.)
    category_id : ItemGroups.Category integer
    description : The text value from the Excel cell (e.g. 'Engineer')
    cache       : Mutable dict passed in from import_excel_process; cleared
                  between requests automatically because it is local to the
                  view call.
    """
    cache_key = (category_id, description.strip().lower())
    if cache_key in cache:
        return cache[cache_key]

    with connections[db_alias].cursor() as cur:
        # ── Look up existing row (case-insensitive) ────────────────────────
        cur.execute(
            'SELECT "GroupID" FROM "ItemGroups" '
            'WHERE "Category" = %s AND LOWER("Description") = LOWER(%s) '
            'LIMIT 1',
            [category_id, description.strip()],
        )
        row = cur.fetchone()
        if row:
            group_id = row[0]
            cache[cache_key] = group_id
            return group_id

        # ── Insert new row ─────────────────────────────────────────────────
        cur.execute('SELECT COALESCE(MAX("GroupID"), 0) FROM "ItemGroups"')
        next_id = cur.fetchone()[0] + 1

        cur.execute(
            'INSERT INTO "ItemGroups" '
            '("GroupID", "Category", "Description", "Under", "NREC") '
            'VALUES (%s, %s, %s, %s, %s)',
            [next_id, category_id, description.strip(), 1, None],
        )
        logger.info(
            '_resolve_itemgroup_id: inserted "%s" → GroupID=%d (Category=%d)',
            description.strip(), next_id, category_id,
        )

    cache[cache_key] = next_id
    return next_id


# ── Process endpoint (shared across all modules) ──────────────────────────────

@require_http_methods(['POST'])
def import_excel_process(request):
    """
    POST body (JSON):
    {
        "import_type": "employee",
        "rows": [ {"Reg No": "001", "Emp Name": "John", ...}, ... ]
    }
    Returns:
    {
        "success": true,
        "imported": 12,
        "errors": [{"row": 3, "error": "..."}, ...]
    }

    ItemGroup columns (flagged with 'itemgroup_col': True in the column config)
    are resolved text → GroupID before the row is coerced, so the final DB
    write always receives an integer.  New ItemGroup rows are auto-inserted when
    the description is not found in the given category.
    """
    try:
        body        = json.loads(request.body)
        import_type = body.get('import_type', '').strip()
        rows        = body.get('rows', [])
    except (json.JSONDecodeError, KeyError):
        return JsonResponse({'success': False, 'error': 'Invalid JSON payload'})

    if import_type not in _IMPORT_REGISTRY:
        return JsonResponse({'success': False, 'error': f'Unknown import type: "{import_type}"'})
    if not isinstance(rows, list) or not rows:
        return JsonResponse({'success': False, 'error': 'No rows to import'})

    config   = _IMPORT_REGISTRY[import_type]
    columns  = config['columns']
    table    = config['table']
    pk_col   = config['pk_col']
    db_alias = config['db_alias_fn']()

    creator = config.get('table_creator')
    if creator:
        try:
            creator(db_alias)
        except Exception as exc:
            logger.warning('import_excel_process: table_creator failed: %s', exc)

    # ── Identify ItemGroup columns once per request ────────────────────────────
    # Build a quick lookup: excel_key → category_id for all itemgroup_col fields
    itemgroup_cols = {
        col['excel']: col['category_id']
        for col in columns
        if col.get('itemgroup_col') and col.get('category_id')
    }

    # Per-request ItemGroup resolution cache — avoids repeated DB hits for the
    # same (category_id, description) pair across many rows.
    ig_cache: dict = {}

    imported = 0
    errors   = []

    for idx, row in enumerate(rows):
        excel_row_num = idx + 2   # Excel row 1 = header
        try:
            # ── Resolve ItemGroup text values → GroupID integers ───────────
            # Mutate a copy so the original row dict is not changed (helpful
            # if the caller retries on error).
            resolved_row = dict(row)
            for excel_key, category_id in itemgroup_cols.items():
                raw_val = str(resolved_row.get(excel_key) or '').strip()
                if not raw_val:
                    # Empty → leave as-is; _coerce_row will handle required check
                    continue
                group_id = _resolve_itemgroup_id(
                    db_alias, category_id, raw_val, ig_cache
                )
                # Overwrite the text with the integer so _coerce_row treats it
                # as a plain 'int' column from this point on.
                resolved_row[excel_key] = str(group_id)

            db_row = _coerce_row(resolved_row, columns)
            _upsert_row(db_alias, table, pk_col, db_row)
            imported += 1
        except Exception as exc:
            errors.append({'row': excel_row_num, 'error': str(exc)})
            logger.warning('import row %d error: %s', excel_row_num, exc)

    return JsonResponse({'success': True, 'imported': imported, 'errors': errors})


# ── Date format list ───────────────────────────────────────────────────────────

_DATE_FORMATS = [
    '%Y-%m-%d',
    '%d/%m/%Y',
    '%d-%m-%Y',
    '%d.%m.%Y',
    '%d/%m/%y',
    '%d-%m-%y',
    '%d.%m.%y',
    '%m/%d/%Y',
    '%m-%d-%Y',
    '%m/%d/%y',
    '%m-%d-%y',
    '%d %b %Y',
    '%d %B %Y',
    '%b %d, %Y',
    '%B %d, %Y',
    '%Y/%m/%d',
]


def _parse_date(excel_key: str, raw: str) -> str:
    """
    Parse a date string from any format Excel or the system datepicker might
    produce and return an ISO-8601 string (YYYY-MM-DD).

    Also handles Excel numeric serial dates (e.g. "45273", "45273.0").
    Excel's epoch is 1899-12-30.

    Raises ValueError with a user-friendly message if parsing fails.
    """
    # ── Excel serial number ───────────────────────────────────────────────
    stripped = raw.split('.')[0]
    if stripped.lstrip('-').isdigit() and not raw.startswith('0') and len(stripped) >= 4:
        serial = int(stripped)
        if 1 <= serial <= 2_958_465:
            try:
                return (date(1899, 12, 30) + timedelta(days=serial)).isoformat()
            except OverflowError:
                pass

    # ── String-format parsing ─────────────────────────────────────────────
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            continue

    raise ValueError(
        f'"{excel_key}": unrecognised date "{raw}". '
        f'Accepted formats: DD/MM/YYYY, DD/MM/YY, YYYY-MM-DD.'
    )


# ── Row helpers ───────────────────────────────────────────────────────────────

def _coerce_row(row: dict, columns: list) -> dict:
    result = {}
    for col in columns:
        excel_key = col['excel']
        db_col    = col['db_col']
        typ       = col.get('type', 'str')
        required  = col.get('required', False)

        raw = str(row.get(excel_key) or '').strip()

        if not raw:
            if required:
                raise ValueError(f'Column "{excel_key}" is required')
            result[db_col] = None
            continue

        if typ == 'int':
            try:
                result[db_col] = int(float(raw))
            except ValueError:
                raise ValueError(f'"{excel_key}": expected integer, got "{raw}"')

        elif typ == 'float':
            try:
                result[db_col] = float(raw)
            except ValueError:
                raise ValueError(f'"{excel_key}": expected number, got "{raw}"')

        elif typ == 'date':
            result[db_col] = _parse_date(excel_key, raw)

        else:
            result[db_col] = raw

    return result


def _upsert_row(db_alias: str, table: str, pk_col: str, data: dict):
    pk_val = data.get(pk_col)
    if not pk_val:
        raise ValueError(f'Primary key "{pk_col}" is missing or empty')

    q = lambda c: f'"{c}"'

    with connections[db_alias].cursor() as cur:
        cur.execute(
            f'SELECT 1 FROM {q(table)} WHERE {q(pk_col)} = %s LIMIT 1',
            [pk_val],
        )
        exists = cur.fetchone() is not None

    cols = [c for c, v in data.items() if v is not None]
    vals = [data[c] for c in cols]

    with connections[db_alias].cursor() as cur:
        if exists:
            update_cols = [c for c in cols if c != pk_col]
            if not update_cols:
                return
            set_clause = ', '.join(f'{q(c)} = %s' for c in update_cols)
            cur.execute(
                f'UPDATE {q(table)} SET {set_clause} WHERE {q(pk_col)} = %s',
                [data[c] for c in update_cols] + [pk_val],
            )
        else:
            col_clause = ', '.join(q(c) for c in cols)
            val_clause = ', '.join('%s' for _ in cols)
            cur.execute(
                f'INSERT INTO {q(table)} ({col_clause}) VALUES ({val_clause})',
                vals,
            )
            